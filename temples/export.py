"""Write the master list to output/south_india_temples.xlsx."""
from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill

from .common import OUTPUT, STATES

XLSX = OUTPUT / "south_india_temples.xlsx"

COLUMNS = [  # (field, header, width)
    ("temple_id", "Temple ID", 16), ("name", "Temple Name", 38), ("name_local", "Local Name", 30),
    ("deity", "Deity (Moolavar)", 24), ("category", "Category / Tradition", 18),
    ("deity_basis", "Deity From", 16),
    ("special_group", "Special Group", 24), ("religion", "Religion", 10),
    ("state", "State", 15), ("district", "District", 18), ("city", "City / Village", 18),
    ("address", "Address", 40), ("pincode", "PIN", 9),
    ("latitude", "Latitude", 11), ("longitude", "Longitude", 11), ("google_maps_url", "Google Maps", 13),
    ("operator", "Managed By", 22), ("phone", "Phone", 15), ("website", "Website", 26),
    ("opening_hours", "Timings", 20), ("heritage", "Heritage Status", 22),
    ("inscription_count", "Inscriptions", 11), ("inscription_refs", "Inscription Refs", 30), ("inception", "Built / Inception", 12),
    ("description", "Description", 36), ("image", "Image", 20),
    ("osm_id", "OSM ID", 16), ("wikidata_id", "Wikidata ID", 12), ("google_place_id", "Google Place ID", 20),
    ("govt_id", "Govt ID", 12), ("sources", "Sources", 30), ("source_count", "No. of Sources", 9),
    ("last_updated", "Last Updated", 12),
]
HEAD_FILL = PatternFill("solid", fgColor="8B1E1E")
HEAD_FONT = Font(bold=True, color="FFFFFF")
LINK_FONT = Font(color="0563C1", underline="single")


def _header(ws, labels):
    cells = []
    for lab in labels:
        c = WriteOnlyCell(ws, value=lab)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
        cells.append(c)
    ws.append(cells)


def _temple_sheet(wb, title, records):
    ws = wb.create_sheet(title[:31])
    for i, (_, _, w) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[_col(i)].width = w
    ws.freeze_panes = "C2"
    _header(ws, [h for _, h, _ in COLUMNS])
    for r in records:
        row = []
        for f, _, _ in COLUMNS:
            v = r.get(f, "")
            if f in ("latitude", "longitude", "source_count") and v != "":
                v = float(v) if f != "source_count" else int(v)
            if f == "google_maps_url" and v:
                c = WriteOnlyCell(ws, value=f'=HYPERLINK("{v}","Open map")')
                c.font = LINK_FONT
                row.append(c)
                continue
            row.append(v)
        ws.append(row)
    ws.auto_filter.ref = f"A1:{_col(len(COLUMNS))}{len(records) + 1}"


def _col(n: int) -> str:
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def _summary(wb, records):
    ws = wb.create_sheet("Summary", 0)
    ws.column_dimensions["A"].width = 34
    for c in "BCDEFGHIJ":
        ws.column_dimensions[c].width = 14
    bold = Font(bold=True, size=13)

    def title(text):
        c = WriteOnlyCell(ws, value=text)
        c.font = bold
        ws.append([])
        ws.append([c])

    t = WriteOnlyCell(ws, value="South India Temples — GPS Directory")
    t.font = Font(bold=True, size=16, color="8B1E1E")
    ws.append([t])
    ws.append([f"Updated {datetime.now():%d %b %Y %H:%M}", f"{len(records):,} temples"])

    title("By state")
    _header(ws, ["State", "Temples", "With GPS", "Named", "Deity known", "Multi-source"])
    by_state = defaultdict(list)
    for r in records:
        by_state[r["state"] or "(unknown)"].append(r)
    for st in list(STATES.values()) + sorted(set(by_state) - set(STATES.values())):
        rs = by_state.get(st, [])
        ws.append([st, len(rs), sum(r["latitude"] != "" for r in rs),
                   sum(r["name"] != "(unnamed temple)" for r in rs),
                   sum(bool(r["deity"]) for r in rs), sum(int(r["source_count"] or 0) > 1 for r in rs)])

    title("By category (tradition)")
    cats = Counter(r["category"] or "(not identified)" for r in records)
    states = list(STATES.values())
    _header(ws, ["Category"] + states + ["Total"])
    for cat, n in cats.most_common():
        ws.append([cat] + [sum(1 for r in by_state.get(s, []) if (r["category"] or "(not identified)") == cat)
                           for s in states] + [n])

    title("Top deities")
    _header(ws, ["Deity", "Temples"])
    for d, n in Counter(r["deity"] for r in records if r["deity"]).most_common(25):
        ws.append([d, n])

    title("Special groups")
    _header(ws, ["Group", "Temples"])
    groups = Counter(g for r in records for g in r["special_group"].split("; ") if g)
    for g, n in groups.most_common():
        ws.append([g, n])

    title("Records per source")
    _header(ws, ["Source", "Temples"])
    srcs = Counter(s.split(":")[0].strip() for r in records for s in set(r["sources"].split("; ")) if s)
    for s, n in srcs.most_common():
        ws.append([s, n])


def _districts(wb, records):
    ws = wb.create_sheet("By District", 1)
    for c, w in zip("ABCDEFGH", (16, 22, 10, 10, 10, 12, 12, 30)):
        ws.column_dimensions[c].width = w
    ws.freeze_panes = "C2"
    _header(ws, ["State", "District", "Temples", "Named", "Deity known", "Village known",
                 "Multi-source", "Most common deity"])
    groups = defaultdict(list)
    for r in records:
        groups[(r["state"] or "(unknown)", r["district"] or "(unknown)")].append(r)
    for (st, d), rs in sorted(groups.items()):
        top = Counter(r["deity"] for r in rs if r["deity"]).most_common(1)
        ws.append([st, d, len(rs), sum(r["name"] != "(unnamed temple)" for r in rs),
                   sum(bool(r["deity"]) for r in rs), sum(bool(r["city"]) for r in rs),
                   sum(int(r["source_count"] or 0) > 1 for r in rs), top[0][0] if top else ""])
    ws.auto_filter.ref = f"A1:H{len(groups) + 1}"


INS_COLUMNS = [
    ("inscription_id", "Inscription ID", 16), ("title", "Title", 34), ("type", "Type", 20),
    ("temple_name", "Temple", 32), ("temple_id", "Temple ID", 16), ("link_method", "How Linked", 24),
    ("location_detail", "Place in Temple", 36), ("village", "Village", 18), ("taluk", "Taluk", 14),
    ("district", "District", 16), ("district_old", "District (as printed)", 16), ("state", "State", 14),
    ("latitude", "Latitude", 11), ("longitude", "Longitude", 11), ("location_basis", "Location Accuracy", 22),
    ("google_maps_url", "Google Maps", 13), ("king", "King", 22), ("dynasty", "Dynasty", 14),
    ("year_ce", "Year (CE)", 9), ("date_text", "Date (as printed)", 26), ("language", "Language", 11),
    ("script", "Script", 11), ("reference", "Reference (A.R. No.)", 20), ("volume", "SII Vol.", 8),
    ("number", "No.", 6), ("summary", "Summary", 50), ("text", "Text", 30), ("image", "Image", 18),
    ("source", "Source", 24), ("source_url", "Source Link", 14),
]


def _inscription_sheet(wb, ins):
    ws = wb.create_sheet("Inscriptions")
    for i, (_, _, w) in enumerate(INS_COLUMNS, start=1):
        ws.column_dimensions[_col(i)].width = w
    ws.freeze_panes = "C2"
    _header(ws, [h for _, h, _ in INS_COLUMNS])
    for r in ins:
        row = []
        for f, _, _ in INS_COLUMNS:
            v = r.get(f, "")
            if f in ("latitude", "longitude") and v != "":
                v = float(v)
            elif f == "year_ce" and str(v).isdigit():
                v = int(v)
            elif f in ("google_maps_url", "source_url") and v:
                c = WriteOnlyCell(ws, value=f'=HYPERLINK("{v}","{"Open map" if f == "google_maps_url" else "Open"}")')
                c.font = LINK_FONT
                row.append(c)
                continue
            row.append(v)
        ws.append(row)
    ws.auto_filter.ref = f"A1:{_col(len(INS_COLUMNS))}{len(ins) + 1}"


def _ins_summary(wb, ins):
    ws = wb.create_sheet("Inscription Summary")
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 12
    for title, key in (("By state", "state"), ("By dynasty", "dynasty"), ("By type", "type"),
                       ("By language", "language"), ("By source", "source"), ("Location accuracy", "location_basis")):
        c = WriteOnlyCell(ws, value=title)
        c.font = Font(bold=True, size=13)
        ws.append([c])
        _header(ws, [title.replace("By ", "").title(), "Inscriptions"])
        for k, n in Counter(r.get(key) or "(not known)" for r in ins).most_common(25):
            ws.append([k, n])
        ws.append([])
    linked = sum(bool(r.get("temple_id")) for r in ins)
    ws.append([f"Linked to a temple: {linked:,} of {len(ins):,}"])
    by_century = Counter((int(r["year_ce"]) // 100 + 1) for r in ins if str(r.get("year_ce", "")).isdigit())
    if by_century:
        c = WriteOnlyCell(ws, value="By century (CE)")
        c.font = Font(bold=True, size=13)
        ws.append([])
        ws.append([c])
        _header(ws, ["Century", "Inscriptions"])
        for cen in sorted(by_century):
            ws.append([f"{cen}th century", by_century[cen]])


def _load_ins():
    from .inscriptions import INS_MASTER
    from .common import read_rows
    return read_rows(INS_MASTER) if INS_MASTER.exists() else []


def _notes(wb):
    ws = wb.create_sheet("About")
    ws.column_dimensions["A"].width = 120
    for line in [
        "Data sources & licences",
        "OpenStreetMap — © OpenStreetMap contributors, ODbL 1.0 (openstreetmap.org/copyright). Attribution required when sharing.",
        "Wikidata — CC0 public domain.",
        "Overture Maps Foundation — Places (CDLA-Permissive-2.0 / Apache-2.0 / CC0 by contributor). Attribution: Overture Maps Foundation.",
        "GeoNames.org — CC BY 4.0. Attribution: GeoNames.org.",
        "ASI — list of Centrally Protected Monuments (Archaeological Survey of India, Government of India).",
        "Google Maps — Places API; place IDs only by default. Follow Google Maps Platform terms for stored content.",
        "Government — files in data/sources/govt (e.g. HR&CE, data.gov.in, ASI). Check each dataset's licence.",
        "Private collection — files in data/sources/private.",
        "",
        "Deity and Category are taken from the source when given, otherwise guessed from the temple name — the 'Deity From' column says which. Verify guessed rows that matter.",
        "City / Village: from the address when given, otherwise the nearest OpenStreetMap village or town (within 4 km).",
        "",
        "Inscriptions",
        "South Indian Inscriptions (SII) volumes, Archaeological Survey of India — index fields (number, A.R. reference, village, district, place in temple, king, date) and a one-line summary, read from the online edition at whatisindia.com. Open the Source Link for the full text.",
        "OpenStreetMap epigraphs / hero stones and Wikidata inscription items — mapped positions.",
        "Your collections — files in data/sources/inscriptions (ARIE lists, TN Archaeology district volumes, field notes).",
        "How Linked / Location Accuracy: 'Exact' = mapped point; 'Temple (matched by name in village)' = SII temple name matched a temple in that village; 'Village (approximate)' = only the village was found — the temple still needs confirming.",
        "District (as printed) is the old district name used in the volume (e.g. Tanjore, South Arcot, Cuddapah).",
        "Category: Shaivam (Shiva), Vaishnavam (Vishnu), Shaktam (Devi/Amman), Kaumaram (Murugan), Ganapatyam (Ganesha), Folk / Grama Devata, Jain, Buddhist.",
        "Google Maps column opens the exact GPS point.",
    ]:
        ws.append([line])


def write_excel(records: list[dict]) -> str:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    wb = Workbook(write_only=True)
    _summary(wb, records)
    _districts(wb, records)
    _temple_sheet(wb, "All Temples", records)
    for st in STATES.values():
        rs = [r for r in records if r["state"] == st]
        if rs:
            _temple_sheet(wb, st, rs)
    ins = _load_ins()
    if ins:
        _ins_summary(wb, ins)
        _inscription_sheet(wb, ins)
    _notes(wb)
    tmp = XLSX.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    tmp.replace(XLSX)
    return str(XLSX)


def write_geojson(records: list[dict]) -> list[str]:
    """One GeoJSON file per state in data/map/ — GitHub shows these as an interactive map."""
    import json
    from pathlib import Path
    out_dir = Path(__file__).resolve().parent.parent / "data" / "map"
    out_dir.mkdir(parents=True, exist_ok=True)
    written = []
    for st in STATES.values():
        feats = []
        for r in records:
            if r["state"] != st or r["latitude"] == "":
                continue
            props = {"name": r["name"], "local name": r["name_local"], "deity": r["deity"],
                     "inscriptions": r.get("inscription_count", ""),
                     "category": r["category"], "village / town": r["city"], "district": r["district"],
                     "special group": r["special_group"], "google maps": r["google_maps_url"],
                     "id": r["temple_id"]}
            feats.append({"type": "Feature",
                          "geometry": {"type": "Point", "coordinates": [float(r["longitude"]), float(r["latitude"])]},
                          "properties": {k: v for k, v in props.items() if v}})
        path = out_dir / f"{st.replace(' ', '_')}.geojson"
        path.write_text(json.dumps({"type": "FeatureCollection", "features": feats},
                                   ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
        written.append(str(path))
    ins = _load_ins()
    for st in STATES.values():
        feats = [{"type": "Feature",
                  "geometry": {"type": "Point", "coordinates": [float(r["longitude"]), float(r["latitude"])]},
                  "properties": {k: v for k, v in {
                      "title": r["title"], "type": r["type"], "temple": r["temple_name"], "king": r["king"],
                      "dynasty": r["dynasty"], "year": r["year_ce"], "reference": r["reference"],
                      "accuracy": r["location_basis"], "source": r["source_url"]}.items() if v}}
                 for r in ins if r.get("state") == st and str(r.get("latitude", "")).strip()]
        if feats:
            path = out_dir / f"Inscriptions_{st.replace(' ', '_')}.geojson"
            path.write_text(json.dumps({"type": "FeatureCollection", "features": feats},
                                       ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
            written.append(str(path))
    return written
